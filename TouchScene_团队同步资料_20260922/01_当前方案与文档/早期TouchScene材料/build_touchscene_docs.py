from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

BLACK = "000000"
NAVY = "17324D"
BLUE = "2D6A9F"
PALE_BLUE = "EDF4FA"
PALE_GRAY = "F6F7F8"
BORDER = "D9D9D9"
GRAY = "5F6B76"


PRODUCT_REFS = [
    ("P1", "WHO", "Blindness and vision impairment", "2026", "全球视力损伤规模与影响背景。", "https://www.who.int/news-room/fact-sheets/detail/blindness-and-visual-impairment"),
    ("P2", "Braille Authority of North America", "Guidelines and Standards for Tactile Graphics", "2022", "触觉图形的简化、留白、层级与可读性原则。", "https://www.brailleauthority.org/guidelines-and-standards-tactile-graphics"),
    ("P3", "Tekli J et al.", "Evaluating touch-screen vibration modality for blind users to access simple shapes and graphics", "2018", "盲人能够通过振动触屏识别简单轮廓图形，同时揭示识别时间与探索覆盖不足等限制。", "https://doi.org/10.1016/j.ijhcs.2017.10.009"),
    ("P4", "Hahn M E et al.", "The Comprehension of STEM Graphics via a Multisensory Tablet Electronic Device by Students with Visual Impairments", "2019", "验证振动与声音结合的触屏设备用于 STEM 图形理解的可行性。", "https://doi.org/10.1177/0145482X19876463"),
    ("P5", "Kristjansson A et al.", "Designing sensory-substitution devices Principles pitfalls and potential", "2016", "指出训练、感官带宽和信息过载是感官替代产品设计的关键约束。", "https://pmc.ncbi.nlm.nih.gov/articles/PMC5044782/"),
    ("P6", "American Printing House", "Monarch", "2026", "商业化多行盲文与触觉图形设备，官方标价和规格用于竞品比较。", "https://www.aph.org/product/monarch/"),
    ("P7", "Dot Inc", "Dot Pad", "2026", "300 单元触觉图形设备及官方价格区间，用于说明专业设备能力与成本。", "https://www.dotincorp.com/en/product/pad"),
    ("P8", "Orbit Research", "Graphiti Interactive Tactile Graphics Display", "2026", "60×40 独立可变高度触点阵列，代表专业动态触觉图形路线。", "https://www.orbitresearch.com/products/blindness-products/tactile-graphic-displays/graphiti-a-breakthrough-in-non-visual-access-to-all-forms-of-graphical-information/"),
    ("P9", "Insta360", "SDK Guide", "2026", "说明 Camera SDK 与 Media SDK 的能力、支持平台和申请方式。", "https://onlinemanual.insta360.com/developer/en-us/resource/sdk"),
]


TECH_REFS = [
    ("R1", "Bach-y-Rita P et al.", "Vision substitution by tactile image projection", "1969", "视觉信息向触觉编码的奠基性研究。", "https://doi.org/10.1038/221963a0"),
    ("R2", "Bach-y-Rita P", "Tactile vision substitution past and future", "1983", "总结视觉转触觉的可行性与触觉通道分辨率限制。", "https://pubmed.ncbi.nlm.nih.gov/6874260/"),
    ("R3", "Kristjansson A et al.", "Designing sensory-substitution devices Principles pitfalls and potential", "2016", "强调训练、注意力容量、带宽差异和避免信息过载。", "https://pmc.ncbi.nlm.nih.gov/articles/PMC5044782/"),
    ("R4", "Tekli J et al.", "Evaluating touch-screen vibration modality for blind users to access simple shapes and graphics", "2018", "为轮廓振动、简单图形和主动探索路线提供直接用户研究依据。", "https://doi.org/10.1016/j.ijhcs.2017.10.009"),
    ("R5", "Hahn M E et al.", "The Comprehension of STEM Graphics via a Multisensory Tablet Electronic Device by Students with Visual Impairments", "2019", "支持振动与语音组合用于图表、地图和 STEM 内容。", "https://doi.org/10.1177/0145482X19876463"),
    ("R6", "Gori M et al.", "Haptic-assistive technologies for audition and vision sensory disabilities", "2018", "系统综述触觉辅助技术，并提出小型化、低成本与个人设备协同方向。", "https://pubmed.ncbi.nlm.nih.gov/29017361/"),
    ("R7", "Braille Authority of North America", "Guidelines and Standards for Tactile Graphics", "2022", "支持简化、删除装饰、控制纹理数量、保留参照点等算法原则。", "https://www.brailleauthority.org/guidelines-and-standards-tactile-graphics"),
    ("R8", "Holloway L et al.", "Refreshable Tactile Displays for Accessible Data Visualisation", "2024", "总结动态触觉图形的机会以及尺寸、触点高度和保真度差异。", "https://arxiv.org/abs/2401.15836"),
    ("R9", "World Wide Web Consortium", "Vibration API", "2026", "定义 navigator.vibrate 及其用户激活和能力边界。", "https://www.w3.org/TR/vibration/"),
    ("R10", "Android Developers", "Implement haptics on Android", "2026", "Android 触觉 API、设备能力检查与预定义效果。", "https://developer.android.com/develop/ui/views/haptics"),
    ("R11", "Android Developers", "Android haptics API reference", "2026", "说明幅度控制并非所有设备支持，支持按节奏降级。", "https://developer.android.com/develop/ui/views/haptics/haptics-apis"),
    ("R12", "MDN Web Docs", "BluetoothRemoteGATTCharacteristic writeValueWithResponse", "2025", "Web Bluetooth 写入确认、HTTPS 限制和浏览器兼容性。", "https://developer.mozilla.org/en-US/docs/Web/API/BluetoothRemoteGATTCharacteristic/writeValueWithResponse"),
    ("R13", "Google Chrome", "Web Bluetooth Notifications Sample", "2026", "GATT 通知订阅的官方示例。", "https://googlechrome.github.io/samples/web-bluetooth/notifications.html"),
    ("R14", "Espressif", "ESP32 C3 ESP IDF Programming Guide", "2026", "ESP32-C3 Bluetooth LE 协议栈与 NimBLE 支持。", "https://docs.espressif.com/projects/esp-idf/en/stable/esp32c3/index.html"),
    ("R15", "Insta360", "SDK Guide", "2026", "Camera SDK 与 Media SDK 能力和支持平台。", "https://onlinemanual.insta360.com/developer/en-us/resource/sdk"),
    ("R16", "Insta360", "Integration Guide", "2026", "SDK 与 OSC 的连接方式、能力和稳定性差异。", "https://onlinemanual.insta360.com/developer/en-us/resource/integration"),
    ("R17", "OpenCV", "Canny Edge Detection", "2026", "边缘检测的降噪、梯度和双阈值处理依据。", "https://docs.opencv.org/4.x/da/d22/tutorial_py_canny.html"),
    ("R18", "OpenCV", "Morphological Transformations", "2026", "开闭运算、腐蚀和膨胀用于清除噪点与连接轮廓。", "https://docs.opencv.org/4.x/d9/d61/tutorial_py_morphological_ops.html"),
]


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_keep(paragraph, keep_next=False, page_break_before=False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if page_break_before:
        p_pr.append(OxmlElement("w:pageBreakBefore"))


def add_hyperlink(paragraph, text: str, url: str, color=BLUE) -> None:
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([r_fonts, color_el, underline])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([r_pr, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    paragraph.add_run(" 页")


def apply_run_font(run, east="Microsoft YaHei", latin="Arial", size=None, bold=None, color=None) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_document(title: str, subject: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.8)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style.font.size = Pt(27)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(18)
    title_p_pr = title_style.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (("Heading 1", 17, 18, 7), ("Heading 2", 13, 13, 5), ("Heading 3", 11, 10, 4)):
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for st_name in ("List Bullet", "List Number"):
        st = styles[st_name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(10.6)
        st.paragraph_format.space_after = Pt(3)

    core = doc.core_properties
    core.title = title
    core.subject = subject
    core.author = "TouchScene 项目组"
    core.keywords = "视觉转触觉, 触觉图形, 无障碍, Insta360, ESP32"

    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.style = styles["Normal"]
    footer.add_run(title + "   ")
    add_page_number(footer)
    for run in footer.runs:
        apply_run_font(run, size=8.5, color=GRAY)
    return doc


def add_cover(doc: Document, title: str, subtitle: str, version: str, audience: str) -> None:
    doc.add_paragraph("TOUCHSCENE", style=None).runs[0].font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(title)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(22)
    run = sub.add_run(subtitle)
    apply_run_font(run, size=14, color=GRAY)
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    widths = [Inches(1.25), Inches(4.9)]
    rows = [("版本", version), ("日期", "2026 年 9 月 18 日"), ("适用读者", audience), ("文档状态", "对外分享版")]
    for i, (k, v) in enumerate(rows):
        meta.cell(i, 0).width = widths[0]
        meta.cell(i, 1).width = widths[1]
        meta.cell(i, 0).text = k
        meta.cell(i, 1).text = v
        set_cell_fill(meta.cell(i, 0), PALE_BLUE)
        for c in meta.rows[i].cells:
            set_cell_margins(c, 105, 130, 105, 130)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for r in para.runs:
                    apply_run_font(r, size=9.5, bold=(c == meta.cell(i, 0)))
    set_table_borders(meta)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run("一句话说明")
    apply_run_font(r, size=10, bold=True, color=GRAY)
    p = doc.add_paragraph()
    r = p.add_run("把相机拍到的简单轮廓和空间关系转换成手指可探索的振动反馈，并用低成本硬件验证从图像到触觉的完整闭环。")
    apply_run_font(r, size=13, bold=True)
    doc.add_page_break()


def add_intro(doc: Document, heading: str, paragraphs: Iterable[str]) -> None:
    doc.add_heading(heading, level=1)
    for text in paragraphs:
        doc.add_paragraph(text)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], font_size=9.2) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.width = Inches(widths[j])
        cell.text = header
        set_cell_fill(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                apply_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for i, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cell = cells[j]
            cell.width = Inches(widths[j])
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if i % 2 == 0:
                set_cell_fill(cell, PALE_GRAY)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(value) > 18 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    apply_run_font(r, size=font_size)
            set_cell_margins(cell)
    # Keep a table header with at least its first body row. Word otherwise may
    # leave a repeated header by itself at the bottom of a page.
    if rows:
        for row in (table.rows[0], table.rows[1]):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    set_repeat_keep(paragraph, keep_next=True)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullets(doc: Document, items: Iterable[str], numbered=False) -> None:
    if numbered:
        for index, item in enumerate(items, start=1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.26)
            p.paragraph_format.first_line_indent = Inches(-0.26)
            p.add_run(f"{index}.  {item}")
        return
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_references(doc: Document, refs) -> None:
    doc.add_heading("参考资料", level=1)
    doc.add_paragraph("以下链接均指向论文原文、标准或官方产品与开发文档。论文用于支撑可行性和设计约束，官方文档用于支撑工程实现。")
    for key, author, title, year, relevance, url in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        r = p.add_run(f"[{key}] {author}. {title}. {year}. ")
        apply_run_font(r, size=9.4)
        add_hyperlink(p, "链接", url)
        r = p.add_run(f"  关联说明：{relevance}")
        apply_run_font(r, size=9.4, color=GRAY)


def build_product_doc() -> Path:
    title = "TouchScene 视觉转触觉产品说明"
    doc = setup_document(title, "面向评委 合作方和产品团队的视觉转触觉产品说明")
    add_cover(doc, title, "从现场影像到可探索触觉反馈的无障碍原型", "1.0", "黑客松评委 合作方 无障碍从业者 产品与研发团队")

    add_intro(doc, "文档摘要", [
        "TouchScene 面向盲人和低视力用户，把相机捕捉到的简单图形、轮廓和空间关系转换为手指可探索的振动反馈。首版以安卓设备为交互面，用户在触觉画布上滑动手指，系统依据当前位置输出背景静默、主体弱振、轮廓强脉冲或关键点双脉冲。",
        "当前原型已经具备本地图片处理、64×48 触觉地图、手机振动、ESP32 BLE 协议、应答校验和自动化测试。实体硬件的真实蓝牙延迟、振动检出率、温升和视障用户识别效果仍需按本文指标实测。产品价值不在于复刻整张照片，而在于快速保留用户真正需要的结构关系。",
    ])

    doc.add_heading("产品解决的问题", level=1)
    doc.add_paragraph("屏幕阅读器擅长文字，却难以直接表达地图、图表、几何图形和物体轮廓。专业动态触觉显示器能够呈现图形，但价格和可得性限制了日常使用。TouchScene 选择手机与低成本振动硬件作为验证入口，让用户在数秒内获得一张可探索的简化触觉图。WHO 估计全球至少有 22 亿人存在近视力或远视力损伤；教育和独立生活中的图形可访问性仍是明确需求。[P1]")
    add_table(doc,
              ["目标用户", "典型任务", "现有障碍", "TouchScene 提供的帮助"],
              [
                  ["盲人学生", "理解几何 轨迹 图表", "文字描述难保留空间关系", "用轮廓和关键点编码结构"],
                  ["低视力用户", "识别现场白板 展品 路线", "取景和细节辨识困难", "先全景捕捉再选取重点区域"],
                  ["教师与讲解员", "快速制作辅助材料", "传统触觉图制作耗时", "照片导入后即时生成简图"],
                  ["无障碍研发团队", "验证触觉交互", "专业触觉硬件成本高", "手机或百元级 BLE 模块快速试验"],
              ], [1.15, 1.55, 1.95, 2.2])

    doc.add_heading("产品定义", level=1)
    doc.add_paragraph("TouchScene 是视觉信息的触觉探索工具，不是盲文阅读器，也不承诺把任意复杂照片完整还原成触觉图。它优先处理边界、包含关系、方向和少量语义锚点。BANA 的触觉图形指南建议对复杂图形做简化、删除非必要装饰、保留参照物并避免过多纹理，这与产品的三层触觉编码一致。[P2]")
    add_table(doc,
              ["触觉状态", "默认反馈", "用户获得的信息"],
              [
                  ["背景", "静默", "这里没有目标内容"],
                  ["主体内部", "低强度短脉冲", "手指位于目标区域内部"],
                  ["外轮廓", "较强而清晰的脉冲", "边界位置和形状变化"],
                  ["关键点", "双脉冲 可配语音", "路口 门口 标签点等语义锚点"],
              ], [1.25, 2.1, 3.3])

    doc.add_heading("核心使用流程", level=1)
    add_bullets(doc, [
        "拍摄或导入图片。首版支持安卓相机拍照和本地图片选择；影石 SDK 获批后接入相机预览、拍摄、文件读取与全景拼接。",
        "选择目标和调节简化程度。系统将图像转换为主体、轮廓和关键点层，并删除不影响任务的细节。",
        "用手指探索。触屏提供绝对坐标，系统查询该位置的触觉状态并立即反馈。",
        "需要更清晰的局部刺激时连接 ESP32 BLE 外设。网页发送坐标和触觉状态，设备执行振动并返回同一序号和坐标。",
        "完成识别任务。用户可以说出形状、路径方向或关键点关系；应用记录用时、正确率和中断情况。",
    ], numbered=True)

    doc.add_heading("为什么需要影石相机", level=1)
    doc.add_paragraph("普通相机要求用户先对准目标。360 度相机可以先覆盖整个现场，再由用户或辅助算法选择白板、地图、展品或路口区域。影石 Camera SDK 提供连接、预览、拍摄和文件管理，Media SDK 提供拼接、预览和导出。桌面端以 USB 为主，Android 支持 USB、Wi-Fi 和蓝牙；SDK 尚未接入时，比赛原型使用影石相机实拍文件导入，保留真实输入链路。[P9]")

    doc.add_heading("首版功能范围", level=1)
    add_table(doc,
              ["能力", "首版状态", "说明"],
              [
                  ["图片导入与本地处理", "已实现", "浏览器本地处理 图片不上传"],
                  ["64×48 触觉地图", "已实现", "背景 主体 轮廓三类已实现 关键点协议已预留"],
                  ["安卓手机振动", "已实现", "依赖浏览器与系统设置 机型间强度不一致"],
                  ["ESP32 BLE 输出", "代码完成", "含命令校验 ACK 队列节流和断连停止"],
                  ["振动传感器确认", "代码完成", "SW-420 只证明振动事件存在 不测振幅质量"],
                  ["影石 SDK 实时接入", "待接入", "先使用实拍照片导入作为可运行路径"],
                  ["视障用户研究", "待执行", "需要真实用户验证辨认率 学习成本和疲劳"],
              ], [1.8, 1.15, 3.75])

    doc.add_heading("优先场景", level=1)
    add_bullets(doc, [
        "课堂中的简单图形，包括几何、函数趋势、轨迹和流程关系。",
        "校园或场馆路线，突出道路、转向、入口和少量地标。",
        "物体外轮廓与构图关系，适合高对比单主体画面。",
        "展览讲解中的局部结构，配合语音标签解释关键点。",
    ])
    doc.add_paragraph("首版不处理文字段落、复杂街景、人物照片细节或持续视频触觉化。这些内容会造成触觉拥挤和认知负担。感官替代研究也提醒设计者控制信息量，并为用户提供训练。[P3][P5]")

    doc.add_heading("方案比较", level=1)
    add_table(doc,
              ["方案", "优势", "限制", "TouchScene 的位置"],
              [
                  ["语音描述", "设备普及 信息密度高", "难同时保留二维空间关系", "与触觉反馈组合使用"],
                  ["静态压凸触觉图", "触感稳定 可双手探索", "制作和分发较慢 不易刷新", "作为教学和演示兜底"],
                  ["专业动态触觉屏", "多点同时呈现 可显示盲文", "设备价格高 首版难自制", "未来可作为高端输出端"],
                  ["手机振动", "零新增硬件 验证最快", "整机振动 局部感弱 机型差异大", "验证交互编码"],
                  ["ESP32 外接振动器", "成本低 可控 可测链路", "仍是单点串行探索", "黑客松主硬件闭环"],
              ], [1.35, 1.7, 2.25, 1.55], font_size=8.8)
    doc.add_paragraph("专业设备展示了长期市场方向：Monarch 具备 3840 个触点，官方标价 15500 美元；Dot Pad 为 300 单元图形显示，官方价格区间为 10000 至 12000 美元；Graphiti 采用 60×40 阵列和可变触点高度。[P6][P7][P8] TouchScene 不以首版硬件替代这些设备，而是用更低门槛验证现场影像到触觉图形的软件链路。")

    doc.add_heading("演示脚本", level=1)
    add_bullets(doc, [
        "影石相机拍摄一个高对比三角形或路线图，并把照片导入 TouchScene。",
        "界面在本地生成触觉地图，展示被保留的主体和轮廓。",
        "评委在安卓触屏上闭眼滑动，先感知背景静默，再追踪轮廓脉冲。",
        "切换 ESP32 BLE 外设，运行四状态自检和 100 点验收，展示坐标 ACK 与可选物理振动确认。",
        "最后说明限制：振动触屏适合简单结构验证，专业触觉屏可作为后续输出端。",
    ], numbered=True)

    doc.add_heading("产品验收指标", level=1)
    add_table(doc,
              ["指标", "原型通过线", "测量方法"],
              [
                  ["图片到触觉图", "不超过 8 秒", "本地计时 20 次取 P95"],
                  ["触点反馈延迟", "平均小于 80 ms P95 小于 150 ms", "BLE 命令到 ACK 往返"],
                  ["坐标映射", "100 点 100% 一致", "序号 模式 X Y 全字段核对"],
                  ["物理振动检出", "非背景至少 99%", "启用 SW-420 时统计状态 3"],
                  ["基础形状识别", "三人平均至少 80%", "圆 三角 矩形 随机顺序闭眼测试"],
                  ["安全与稳定", "10 分钟无断连或异常发热", "连续探索并记录温升和断连"],
              ], [1.35, 2.25, 3.25])

    doc.add_heading("主要风险和应对", level=1)
    add_table(doc,
              ["风险", "影响", "应对"],
              [
                  ["手机整机振动缺少局部性", "边界追踪精度有限", "外接指尖振动器 并使用节奏而非幅度编码"],
                  ["复杂图片产生触觉噪声", "识别失败或疲劳", "限制目标数量 删除纹理 强制预览简化结果"],
                  ["Web Bluetooth 浏览器受限", "部分设备无法连接", "安卓 Chrome 为目标环境 原生 APK 作为后续路线"],
                  ["算法和用户差异", "同一图形对不同人效果不同", "保存个人阈值 提供训练模式 与视障用户共创"],
                  ["过度宣传", "把实验原型误当辅助器具", "对外明确当前状态 不用于安全导航或替代专业训练"],
              ], [1.65, 1.65, 3.55])

    doc.add_heading("迭代路线", level=1)
    add_table(doc,
              ["阶段", "目标", "交付"],
              [
                  ["阶段一 黑客松", "证明图像到触觉闭环", "PWA ESP32 BLE 原型 三类触觉编码 验收数据"],
                  ["阶段二 用户验证", "确认场景和学习成本", "10 至 20 名视障用户任务研究 编码和引导策略修订"],
                  ["阶段三 相机集成", "减少取景和导入步骤", "影石预览 拍摄 后取景 目标选择和离线处理"],
                  ["阶段四 多输出端", "适配不同预算与任务", "手机振动 低成本外设 专业触觉屏和压凸输出适配器"],
              ], [1.45, 2.1, 3.3])

    add_references(doc, PRODUCT_REFS)
    path = OUT / "TouchScene视觉转触觉产品说明.docx"
    doc.save(path)
    return path


def build_technical_doc() -> Path:
    title = "TouchScene 视觉转触觉技术方案"
    doc = setup_document(title, "视觉转触觉原型的系统架构 协议 硬件 测试和参考文献")
    add_cover(doc, title, "安卓触觉地图与 ESP32 BLE 闭环实现", "1.0", "研发团队 硬件工程师 无障碍研究者 技术评审和合作方")

    add_intro(doc, "技术结论", [
        "首版采用安卓触屏绝对坐标加单点触觉反馈。图像被压缩为 64×48 触觉地图；手指移动时，软件只查询当前位置并输出背景、主体、轮廓或关键点状态。该架构避免自制数百个升降触点，同时保留二维探索过程。",
        "输出分为两级。手机振动用于零外设验证；ESP32-C3 通过 Web Bluetooth 驱动外接振动马达，并以 ACK 返回序号、状态和坐标。可选 SW-420 传感器检测实际振动事件，使软件映射正确与物理执行成功能够分别统计。",
        "现有仓库已经完成协议、坐标映射、触觉分类、延迟汇总和浏览器流程的自动化测试。真实手机、ESP32、马达和传感器仍需按验收表完成实测。本方案是研究和演示原型，不应直接用于安全导航。",
    ])

    doc.add_heading("设计依据", level=1)
    doc.add_paragraph("视觉转触觉属于感官替代。早期研究已经证明视觉信息可映射到触觉通道，但触觉带宽、训练成本和设备可用性限制了日常应用。[R1][R2][R3] 本方案不传递完整画面，而是围绕轮廓、主体内部和关键点建立低信息量编码。触屏振动研究显示，盲人能够识别简单轮廓图形，但识别可能较慢，手指探索路径也可能覆盖不足，因此系统必须提供边框、方向标记、训练图和语音提示。[R4][R5]")
    doc.add_paragraph("触觉图形应服务于任务。BANA 指南要求复杂图形简化、消除非必要内容、避免过多纹理，并保留参照点。[R7] 这些原则直接转化为算法约束：限制主体数量、降低纹理、过滤孤点、保持轮廓连续、为关键位置添加少量语义锚点。")

    doc.add_heading("系统边界", level=1)
    add_table(doc,
              ["包含", "不包含"],
              [
                  ["安卓 Chrome PWA 图片导入和相机拍摄", "iOS Safari Web Bluetooth 支持承诺"],
                  ["简单高对比主体和路线的触觉化", "任意照片的完整语义还原"],
                  ["手机振动与单马达 BLE 外设", "首版自制多点升降触觉阵列"],
                  ["坐标和执行事件闭环", "SW-420 对振幅 频率或舒适度的精确测量"],
                  ["离线图像处理和演示验收", "医疗器械 安全导航或替代定向行走训练"],
              ], [3.35, 3.35])

    doc.add_heading("总体架构", level=1)
    add_table(doc,
              ["层级", "组件", "职责", "主要接口"],
              [
                  ["输入层", "安卓相机 本地图片 影石相机", "采集目标图形", "File API Camera SDK Media SDK"],
                  ["处理层", "浏览器图像处理模块", "灰度 阈值 边缘 清理 降采样", "Canvas ImageData 可升级 OpenCV.js"],
                  ["表示层", "64×48 触觉地图", "保存四类触觉状态", "Uint8Array"],
                  ["交互层", "触觉探索画布", "触点绝对坐标映射 节流 最新点队列", "Pointer Events"],
                  ["输出层一", "安卓手机振动", "节奏型触觉验证", "navigator.vibrate"],
                  ["输出层二", "ESP32-C3 外设", "马达控制 看门狗 物理检测", "BLE GATT 写入与通知"],
                  ["验证层", "ACK 与测试面板", "坐标一致性 物理确认 RTT 统计", "序号 X Y mode status"],
              ], [1.0, 1.55, 2.65, 1.65], font_size=8.6)
    doc.add_paragraph("运行序列：图片进入浏览器后先生成触觉地图；手指坐标归一化到网格；应用查询状态并输出振动；BLE 模式下设备回传 ACK；页面分别显示映射是否一致、是否检测到物理振动以及往返延迟。")

    doc.add_heading("图像处理流程", level=1)
    add_table(doc,
              ["步骤", "处理", "目的", "首版实现"],
              [
                  ["一", "缩放和灰度化", "控制计算量并统一输入", "Canvas ImageData"],
                  ["二", "主体阈值", "分离高对比主体与背景", "可调全局阈值"],
                  ["三", "局部梯度", "检测明暗突变形成轮廓", "邻域差分阈值"],
                  ["四", "结构清理", "去除孤点 连接短裂缝", "首版轻量规则 后续 OpenCV 形态学"],
                  ["五", "降采样", "生成 64×48 网格", "按输出网格聚合"],
                  ["六", "状态编码", "背景 主体 轮廓 关键点", "每格 0 至 3"],
              ], [0.55, 1.65, 2.25, 2.4], font_size=8.8)
    doc.add_paragraph("产品化版本建议使用 OpenCV 的自适应阈值、Canny 和形态学开闭运算。Canny 包含降噪、梯度、非极大值抑制和双阈值连接；开闭运算适合去除小噪点和连接轮廓。[R17][R18] 对复杂照片可增加目标分割，但分割结果仍需按触觉规范简化，而不是直接输出所有边缘。")

    doc.add_heading("坐标与触觉编码", level=1)
    doc.add_paragraph("触屏提供绝对坐标，可避免鼠标相对位移长期累积。设触点为 px 和 py，画布左上角为 left 和 top，实际宽高为 width 和 height，网格尺寸为 64×48。映射先将坐标限制在画布边界，再按比例取整到 0 至 63 和 0 至 47。页面尺寸变化后必须重新读取边界，不能缓存旧尺寸。")
    add_table(doc,
              ["模式", "数值", "手机反馈", "ESP32 默认模式", "设计目的"],
              [
                  ["背景", "0", "停止", "马达关闭", "建立清晰基线"],
                  ["主体", "1", "22 ms 轻脉冲", "120 ms 周期内开启 28 ms", "表达区域内部"],
                  ["轮廓", "2", "42 ms 强脉冲", "180 ms 周期内开启 48 ms", "突出边界"],
                  ["关键点", "3", "32 34 32 ms 双脉冲", "360 ms 周期内两次 36 ms", "表达语义锚点"],
              ], [1.0, 0.55, 1.65, 2.15, 1.5], font_size=8.5)
    doc.add_paragraph("网页 Vibration API 只能提供简单节奏，且依赖用户激活和浏览器实现；Android 不同设备对幅度控制的支持也不同，因此首版用节奏区分类别，不把指定幅度视为可靠能力。[R9][R10][R11]")

    doc.add_heading("BLE 服务与数据协议", level=1)
    add_table(doc,
              ["对象", "UUID", "属性"],
              [
                  ["TouchScene 服务", "7b100001-6c7d-4c7a-9a31-54bf3f010001", "Primary Service"],
                  ["命令特征", "7b100002-6c7d-4c7a-9a31-54bf3f010001", "Write With Response"],
                  ["应答特征", "7b100003-6c7d-4c7a-9a31-54bf3f010001", "Notify Read"],
              ], [1.35, 3.65, 1.8], font_size=8.8)
    doc.add_heading("命令包十字节", level=2)
    add_table(doc,
              ["字节", "字段", "说明"],
              [
                  ["0", "magic", "固定 0xA1"], ["1", "seq", "0 至 255 循环序号"], ["2", "mode", "0 至 3"],
                  ["3 至 4", "x", "小端序网格 X"], ["5 至 6", "y", "小端序网格 Y"], ["7", "intensity", "0 至 255 提示值"],
                  ["8", "duration", "10 ms 为单位"], ["9", "checksum", "前九字节 XOR"],
              ], [0.85, 1.35, 4.6], font_size=9.0)
    doc.add_heading("应答包九字节", level=2)
    add_table(doc,
              ["字节", "字段", "说明"],
              [
                  ["0", "magic", "固定 0xA2"], ["1", "seq", "返回命令序号"], ["2", "mode", "已应用模式"], ["3", "status", "执行状态"],
                  ["4 至 5", "x", "返回 X"], ["6 至 7", "y", "返回 Y"], ["8", "checksum", "前八字节 XOR"],
              ], [0.85, 1.35, 4.6], font_size=9.0)
    add_table(doc,
              ["状态", "含义", "页面判定"],
              [
                  ["0", "GPIO 状态已应用 未启用物理传感器", "映射可通过 物理执行未确认"],
                  ["1", "数据包或校验错误", "拒绝"],
                  ["2", "触觉模式无效", "拒绝"],
                  ["3", "传感器窗口内检测到振动", "映射和物理执行均通过"],
                  ["4", "坐标和模式正确 但未检测到振动", "映射通过 物理执行失败"],
              ], [0.75, 3.05, 3.0])
    doc.add_paragraph("Web Bluetooth 需要 HTTPS，连接选择器必须由用户操作触发。writeValueWithResponse 和通知并非所有浏览器都支持，因此首版明确把安卓 Chrome 作为目标环境。[R12][R13]")

    doc.add_heading("队列 超时和安全状态", level=1)
    add_bullets(doc, [
        "同一时刻只允许一个 BLE 写请求在途。手指快速移动时不排队所有点，只保留最新坐标，防止延迟累积。",
        "应用在手指停留时每 180 ms 发送 keepalive。ESP32 在 600 ms 未收到有效命令时关闭马达。",
        "页面抬指、切换到背景、BLE 断连或协议错误时都进入背景状态。ESP32 断连回调直接关闭马达并重新广播。",
        "XOR 只能发现常见传输错误，不能提供安全认证。若产品化涉及敏感数据或不可信环境，应使用平台配对和应用层认证。",
        "触觉反馈频率必须节流，避免持续强振造成不适。用户应始终能够一键停止。",
    ])

    doc.add_heading("硬件设计", level=1)
    doc.add_paragraph("ESP32-C3 负责 BLE 和马达节奏。GPIO 只驱动逻辑电平 MOSFET 的栅极，不能直接给马达供电。5 V 马达使用独立稳压电源，电源地与 ESP32 共地；马达两端并联 SS14 或 1N5819 续流二极管。ESP32-C3 支持 Bluetooth LE，官方 ESP-IDF 提供相应协议栈和示例。[R14]")
    add_table(doc,
              ["部件", "规格", "数量", "作用"],
              [
                  ["ESP32-C3 SuperMini", "Type C 支持 BLE", "1", "控制器"],
                  ["扁平振动马达", "5 V 10 mm ERM", "1", "触觉输出"],
                  ["AO3400 MOSFET 模块", "3.3 V 逻辑电平", "1", "功率开关"],
                  ["续流二极管", "SS14 或 1N5819", "1", "吸收反向电压"],
                  ["栅极电阻与下拉", "100 欧与 10 k 欧", "各 1", "限制尖峰并确保关断"],
                  ["5 V 电源", "稳压 2 A", "1", "马达独立供电"],
                  ["SW-420 可选", "LM393 数字输出", "1", "检测振动事件"],
              ], [1.55, 2.25, 0.7, 2.2], font_size=8.8)
    doc.add_paragraph("仓库参考清单给出的基础闭环物料成本为 36 至 99 元，增加振动传感器后为 39 至 107 元。该价格是原型采购区间，不含外壳、手机、影石相机、运费和人工。")

    doc.add_heading("物理反馈检测", level=1)
    doc.add_paragraph("启用 SW-420 后，固件在启动马达前读取静止电平，兼容高电平触发和低电平触发的模块。中断锁存短脉冲；8 ms 后开始判定，并在 160 ms 窗口内寻找相对基线的变化。检测成功返回状态 3，超时返回状态 4。")
    doc.add_paragraph("SW-420 是开关型传感器，只能证明模块检测到了振动事件。它不能证明振幅足够、频率正确、触感舒适，也不能替代视障用户测试。传感器应固定在马达外壳附近，并调节灵敏度，使马达启动可靠触发而桌面轻触不误触。")

    doc.add_heading("影石相机集成", level=1)
    add_table(doc,
              ["阶段", "实现", "原因"],
              [
                  ["当前可运行路径", "影石相机拍照后导入 JPG", "不依赖 SDK 审批即可完成真实输入闭环"],
                  ["Android 正式集成", "Camera SDK 连接预览拍摄 Media SDK 拼接和后取景", "减少手工导入并支持先拍后选"],
                  ["兼容兜底", "OSC Wi-Fi 控制加本地文件导入", "兼容性较强 但能力和稳定性低于 SDK"],
              ], [1.55, 3.25, 1.95])
    doc.add_paragraph("Insta360 官方资料说明 SDK 由 Camera SDK 和 Media SDK 组成，支持 Windows、Linux、Android 和 iOS。桌面 SDK 主要通过 USB，Android 支持 USB、Wi-Fi 和蓝牙；OSC 走 Wi-Fi，兼容性更强但不包含部分预览、直播和媒体处理能力。[R15][R16]")

    doc.add_heading("代码与验证状态", level=1)
    add_table(doc,
              ["项目", "证据", "状态"],
              [
                  ["触觉地图和协议核心", "dist/touchpuck/core.mjs", "已实现"],
                  ["安卓 PWA 交互", "dist/touchpuck/index.html 和 app.mjs", "已实现"],
                  ["ESP32 BLE 固件", "firmware/touchpuck_ble/touchpuck_ble.ino", "已实现"],
                  ["自动烧录与串口检查", "scripts/touchpuck-commission.ps1", "已实现"],
                  ["核心单元测试", "tests/touchpuck-core.test.mjs", "已通过"],
                  ["浏览器流程测试", "tests/touchpuck-browser-qa.mjs", "已有测试脚本和截图"],
                  ["真实硬件 BLE 与温升", "需手机 ESP32 马达 SW-420", "待实测"],
                  ["真实视障用户任务", "需伦理和招募安排", "待执行"],
              ], [2.0, 3.15, 1.55], font_size=8.7)

    doc.add_heading("测试方案", level=1)
    add_table(doc,
              ["测试", "通过标准", "失败时处理"],
              [
                  ["坐标边界映射", "四角 中心 越界均落入正确网格", "修正边界读取和 clamp"],
                  ["协议往返", "seq mode x y 完全一致", "检查小端序和字段长度"],
                  ["校验故障注入", "任一被破坏数据包必须拒绝", "检查 XOR 和包长"],
                  ["100 点 BLE", "100/100 映射一致", "降低发送速率 排查队列"],
                  ["物理振动", "非背景至少 99/100 返回状态 3", "固定传感器 调阈值 查供电"],
                  ["BLE RTT", "平均小于 80 ms P95 小于 150 ms", "减少在途写入和日志"],
                  ["断连安全", "600 ms 内停止马达", "检查回调和看门狗"],
                  ["十分钟热稳定", "无异常发热和复位", "降低占空比 更换电源或驱动"],
                  ["基础形状识别", "三人平均至少 80%", "简化图形 优化节奏和引导"],
              ], [1.45, 2.9, 2.35], font_size=8.7)

    doc.add_heading("四十八小时实施顺序", level=1)
    add_table(doc,
              ["时间", "工作", "退出条件"],
              [
                  ["0 至 4 小时", "手机振动和内置图形自检", "三种反馈可稳定区分"],
                  ["4 至 10 小时", "ESP32 马达接线 固件烧录", "四状态自检通过"],
                  ["10 至 16 小时", "SW-420 固定和灵敏度校准", "100 次至少 99 次确认"],
                  ["16 至 26 小时", "相机图片 触觉简化 场景素材", "真实照片到触觉图不超过 8 秒"],
                  ["26 至 36 小时", "外壳 走线 交互引导", "连续十分钟稳定"],
                  ["36 至 44 小时", "闭眼内部测试和参数冻结", "圆 三角 矩形达到通过线"],
                  ["44 至 48 小时", "录屏 路演 备份", "离线 Demo 固件 备用图片齐全"],
              ], [1.2, 3.4, 2.1], font_size=8.8)

    doc.add_heading("后续技术路线", level=1)
    add_bullets(doc, [
        "用影石预览和后取景减少用户对准目标的负担，并加入语音选择目标区域。",
        "用 OpenCV.js 或原生 OpenCV 增加透视矫正、自适应阈值、轮廓清理和多尺度简化。",
        "用真实视障用户数据优化节奏、训练流程、边框和关键点密度，不以健全人蒙眼测试代替目标用户研究。",
        "抽象统一输出接口，把同一触觉地图适配到专业动态触觉屏、压凸机或未来锁存点阵。",
        "若单点探索不足以支持目标任务，再评估多点设备。专业动态触觉屏的能力和使用需求可参考现有研究与产品。[R8]",
    ])

    add_references(doc, TECH_REFS)
    path = OUT / "TouchScene视觉转触觉技术方案.docx"
    doc.save(path)
    return path


def main() -> None:
    product = build_product_doc()
    technical = build_technical_doc()
    print(product)
    print(technical)


if __name__ == "__main__":
    main()
